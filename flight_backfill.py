#!/usr/bin/env python3
"""
flight_backfill.py - standalone FR24 -> ADSB DB -> Billing backfill tool.

Input is the operator's Excel/CSV list (REG, CALL SIGN, OPERATOR, TYPE, DEPARTURE, DESTINATION, ROUTE);
the operation date comes from a DATE column in the sheet or, for rows without one, the --date YYYY-MM-DD option.

Commands (run in order, each reads the previous command's state file):

  check   test the ADSB DB and Billing endpoint (also runs automatically before adsb/billing)
  all     check -> fetch -> adsb -> billing -> verify in one go

  fetch   FR24 summary + track for every listed flight  -> state/fr24_flights.json
  adsb    insert path into svl_track_raw, derive track_segment/flight_sector
          with ADSB's own SQL functions, overlay FR24 data -> state/adsb_flights.json
  billing POST the flight_sector rows to Billing's third-party flight-dump API
          (same payload ADSB's Billing.send() produces)   -> state/billing_result.json
  verify  read-only cross-check of all three stages; makes no writes

Config comes from environment / .env (see .env.example).
"""
import argparse
import csv
import json
import logging
import os
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests
from dotenv import load_dotenv

log = logging.getLogger("flight_backfill")
logging.getLogger("fr24sdk").setLevel(logging.ERROR)  # the SDK logs whole HTML error pages
logging.getLogger("httpx").setLevel(logging.WARNING)

FR24_SLEEP_S = 2.0  # FR24 plan budget: 30 requests/minute
GROUND_ALT = -999   # ADSB sentinel for "on the ground" (tracker.py)
NO_ALT = -1000      # ADSB sentinel for "no altitude"

# Same column list/shape as flight_tracking.FT_compile_flights() sends to Billing.
FLIGHT_PAYLOAD_SQL = """
select json_strip_nulls(row_to_json(fs)) from (
    select
        flight_sector_id
        ,flight_sector_id as source_reference_id
        ,icao24
        ,to_timestamp(flight_start_e::int) as flight_start
        ,to_timestamp(flight_end_e::int) as flight_end
        ,mvt_type
        ,nullif(trim(callsign), '') as callsign
        ,mode_a
        ,mode_a_category
        ,aircraft_reg
        ,operator
        ,aircraft_type
        ,origin
        ,destination
        ,route_source
        ,to_timestamp(fir_entry_time_e::int) as fir_entry_time
        ,fir_entry_source
        ,initial_airport_icao
        ,initial_bay_nb
        ,to_timestamp(out_of_blocks_e::int) out_of_blocks_time
        ,takeoff_airport
        ,to_timestamp(takeoff_time_e::int) as takeoff_time
        ,takeoff_source
        ,landing_airport
        ,to_timestamp(landing_time_e::int) as landing_time
        ,landing_source
        ,final_airport_icao
        ,final_bay_nb
        ,to_timestamp(in_blocks_e::int) as in_blocks_time
        ,to_timestamp(fir_exit_time_e::int) as fir_exit_time
        ,fir_exit_source
        ,fir_in_route
        ,fir_in_lat
        ,fir_in_lon
        ,fir_out_route
        ,fir_out_lat
        ,fir_out_lon
        ,flight_state
        ,fr24_call
        ,fr24_response_data
        ,fr24_enriched_at
        ,'adsb' as flight_source
    from flight_tracking.flight_sector
    where flight_sector_id = %s
) fs
"""


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

def env(name, default=None, required=False):
    val = os.environ.get(name, default)
    if required and not val:
        sys.exit(f"Missing required setting {name} (set it in the environment or .env)")
    return val


def parse_ts(value):
    """FR24 ISO-8601 string -> aware UTC datetime."""
    dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


def load_json(path):
    if not path.exists():
        sys.exit(f"State file {path} not found - run the previous command first.")
    return json.loads(path.read_text())


def save_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, default=str))
    log.info("wrote %s", path)


HEADER_ALIASES = {
    "a/c registration": "registration", "registration": "registration", "reg": "registration",
    "call sign": "callsign", "callsign": "callsign",
    "operater": "operator", "operator": "operator",
    "a/c type": "type", "type": "type",
    "departure": "departure", "destination": "destination",
    "route": "route", "date": "date", "fr24_id": "fr24_id",
}


def _sheet_rows(path, sheet=None):
    """Yield rows as lists of cell strings from .xlsx / .csv / .tsv."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        if sheet and sheet not in wb.sheetnames:
            sys.exit(f"{path}: no tab {sheet!r}; tabs are {wb.sheetnames}")
        ws = wb[sheet] if sheet else wb.active
        for row in ws.iter_rows(values_only=True):
            yield ["" if c is None else str(c).strip() for c in row]
    else:
        with open(path, newline="") as fh:
            delim = "\t" if suffix in (".tsv", ".txt") else ","
            for row in csv.reader(fh, delimiter=delim):
                yield [c.strip() for c in row]


def read_flight_list(path, default_date=None, sheet=None):
    rows = _sheet_rows(path, sheet)
    header = None
    for row in rows:  # first row containing a call sign header
        mapped = [HEADER_ALIASES.get(c.lower(), "") for c in row]
        if "callsign" in mapped:
            header = mapped
            break
    if header is None:
        sys.exit(f"{path}: no header row with a 'CALL SIGN' column found")
    flights = []
    for row in rows:
        rec = {h: row[i] for i, h in enumerate(header) if h and i < len(row)}
        if not rec.get("callsign"):
            continue
        day = rec.get("date") or default_date
        if not day:
            sys.exit(f"{path}: no date for {rec['callsign']} - add a DATE column or pass --date YYYY-MM-DD")
        day = day[:10]  # tolerate "2026-09-27 00:00:00" from Excel
        try:
            datetime.strptime(day, "%Y-%m-%d")
        except ValueError:
            sys.exit(f"{path}: bad date {day!r} for {rec['callsign']}, expected YYYY-MM-DD")
        flights.append({
            "callsign": rec["callsign"].replace(" ", "").upper(),
            "date": day,
            "fr24_id": rec.get("fr24_id") or None,
            "registration": (rec.get("registration") or "").upper() or None,
            "operator": rec.get("operator") or None,
            "type": (rec.get("type") or "").upper() or None,
            "departure": (rec.get("departure") or "").upper() or None,
            "destination": (rec.get("destination") or "").upper() or None,
            "route": rec.get("route") or None,
        })
    return flights


def flight_key(item):
    return f"{item['callsign']}@{item['date']}"


def pick_best(matches, hint=None):
    """Rank candidates: sheet registration/route match first, then finished flight, then recency."""
    hint = hint or {}

    def recency(m):
        for k in ("last_seen", "datetime_landed", "datetime_takeoff", "first_seen"):
            if m.get(k):
                return parse_ts(m[k])
        return datetime.min.replace(tzinfo=timezone.utc)

    def score(m):
        eq = lambda a, b: bool(a and b and a.upper() == b.upper())
        return (eq(m.get("reg"), hint.get("registration")),
                eq(m.get("orig_icao"), hint.get("departure")) + eq(m.get("dest_icao"), hint.get("destination")),
                m.get("flight_ended") is True,
                recency(m))
    return max(matches, key=score)


_FIR = None


def load_fir():
    """South Sudan FIR polygon - same ssd_airspace boundary ADSB's GeoFilter uses."""
    global _FIR
    if _FIR is None:
        from shapely.geometry import shape
        from shapely.ops import unary_union
        from shapely.prepared import prep
        path = Path(env("FIR_GEOJSON", str(Path(__file__).with_name("ssd_airspace.geojson"))))
        data = json.loads(path.read_text())
        poly = unary_union([shape(f["geometry"]) for f in data["features"]])
        _FIR = (poly, prep(poly))
    return _FIR


def _interp_alt(a, b, frac):
    return round(a + (b - a) * frac) if a > 0 and b > 0 else (a if frac < 0.5 else b)


def build_path_points(tracks):
    """FR24 track points -> [(lon, lat, alt, epoch)] clipped to the South Sudan FIR.

    The path is clipped geometrically: real points inside the FIR are kept, and where a
    segment between two FR24 points crosses the FIR boundary (including across a gap in
    FR24 coverage) the boundary crossing is added, with time and altitude interpolated.
    ADSB conventions: alt in feet, -999 on the ground, -1000 when unknown; M is epoch seconds.
    """
    from shapely.geometry import LineString, MultiPoint, Point
    poly, fir = load_fir()
    boundary = poly.boundary

    raw, last = [], None
    for t in sorted(tracks, key=lambda p: parse_ts(p["timestamp"])):
        epoch = parse_ts(t["timestamp"]).timestamp()
        if last is not None and epoch <= last:
            continue
        alt = t.get("alt")
        alt = NO_ALT if alt is None else (GROUND_ALT if alt <= 0 else int(alt))
        raw.append((float(t["lon"]), float(t["lat"]), alt, epoch))
        last = epoch

    out = []

    def add(pt):
        if not out or pt[3] > out[-1][3]:
            out.append(pt)

    for i, a in enumerate(raw):
        a_in = fir.covers(Point(a[0], a[1]))
        if a_in:
            add(a)
        if i + 1 == len(raw):
            break
        b = raw[i + 1]
        b_in = fir.covers(Point(b[0], b[1]))
        if a_in and b_in:
            continue
        seg = LineString([(a[0], a[1]), (b[0], b[1])])
        hit = seg.intersection(boundary)
        if hit.is_empty:
            continue
        pts = [hit] if hit.geom_type == "Point" else [g for g in getattr(hit, "geoms", []) if g.geom_type == "Point"]
        total = seg.length or 1.0
        for c in sorted(pts, key=lambda g: seg.project(g)):
            frac = seg.project(c) / total
            add((c.x, c.y, _interp_alt(a[2], b[2], frac), a[3] + (b[3] - a[3]) * frac))
    return out


def path_ewkt(pts):
    body = ", ".join(f"{lon:.7f} {lat:.7f} {alt} {m:.3f}" for lon, lat, alt, m in pts)
    return f"SRID=4326;LINESTRING ZM({body})"


def pg_connect():
    import psycopg
    return psycopg.connect(
        host=env("ADSB_DB_HOST", "localhost"),
        port=env("ADSB_DB_PORT", "5432"),
        dbname=env("ADSB_DB_NAME", required=True),
        user=env("ADSB_DB_USER", required=True),
        password=env("ADSB_DB_PASSWORD", ""),
        connect_timeout=30,
    )


# --------------------------------------------------------------------------
# preflight checks
# --------------------------------------------------------------------------

class PreflightError(Exception):
    pass


def check_adsb():
    """ADSB database must accept a connection and have the flight_tracking tables we write to."""
    try:
        with pg_connect() as conn:
            conn.execute("select 1")
            missing = [t for t in ("svl_track_raw", "svl_flight_raw", "track_segment", "flight_sector", "process_control")
                       if not conn.execute("select to_regclass(%s)", (f"flight_tracking.{t}",)).fetchone()[0]]
            if missing:
                raise PreflightError(f"ADSB database reachable but tables missing in flight_tracking: {missing}")
            host, db = conn.info.host, conn.info.dbname
    except PreflightError:
        raise
    except SystemExit:
        raise
    except Exception as exc:
        raise PreflightError(f"ADSB database connection failed: {exc}") from exc
    log.info("preflight OK: ADSB database %s/%s", host, db)


def check_billing():
    """Billing must be reachable, accept our static token, and accept calls on the flight-dump endpoint.

    Sends an empty list to the dump endpoint: it authenticates and passes the IP/company
    permissions, processes zero flights, and writes nothing."""
    root = env("BILLING_API_ROOT", required=True).rstrip("/")
    try:
        token = billing_token(root)
    except Exception as exc:
        raise PreflightError(f"Billing token endpoint {root}/api/third-party/token/ failed: {exc}") from exc
    try:
        r = requests.post(f"{root}/api/third-party-flight-dump/",
                          headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
                          json=[], timeout=15)
    except Exception as exc:
        raise PreflightError(f"Billing flight-dump endpoint unreachable: {exc}") from exc
    if r.status_code in (401, 403):
        raise PreflightError(f"Billing flight-dump endpoint rejected the token/IP (HTTP {r.status_code}): {r.text[:200]}")
    if r.status_code >= 500:
        raise PreflightError(f"Billing flight-dump endpoint server error (HTTP {r.status_code}): {r.text[:200]}")
    log.info("preflight OK: Billing %s (token accepted, flight-dump endpoint answered HTTP %s)", root, r.status_code)


def run_checks(*names):
    checks = {"adsb": check_adsb, "billing": check_billing}
    failed = []
    for n in names:
        try:
            checks[n]()
        except PreflightError as exc:
            log.error("preflight FAILED: %s", exc)
            failed.append(n)
    if failed:
        log.error("cancelled: preflight check(s) failed: %s - nothing was run", ", ".join(failed))
        return False
    return True


def cmd_check(args):
    return 0 if run_checks("adsb", "billing") else 1


# --------------------------------------------------------------------------
# 1. fetch
# --------------------------------------------------------------------------

def fr24_call(fn, *a, **kw):
    """Call the FR24 SDK, retrying transient failures (502/503/504, 429, network errors)."""
    delays = (5, 15, 30)
    for attempt in range(len(delays) + 1):
        try:
            return fn(*a, **kw)
        except Exception as exc:
            msg = str(exc)
            transient = any(t in msg for t in ("502", "503", "504", "429", "Bad Gateway", "Timeout", "timed out",
                                               "Connection", "Too Many Requests"))
            if not transient or attempt == len(delays):
                raise RuntimeError(msg.split("\n")[0][:200]) from exc
            log.warning("FR24 transient error (%s) - retry %d/%d in %ss", msg.split("\n")[0][:80],
                        attempt + 1, len(delays), delays[attempt])
            time.sleep(delays[attempt])


def cmd_fetch(args):
    from fr24sdk.client import Client

    flights = read_flight_list(args.input, args.date, args.sheet)
    out, failures = [], 0
    with Client(api_token=env("FR24_API_TOKEN", required=True)) as client:
        for item in flights:
            key = flight_key(item)
            day = datetime.strptime(item["date"], "%Y-%m-%d")
            try:
                if item["fr24_id"]:
                    resp = fr24_call(client.flight_summary.get_full, flight_ids=[item["fr24_id"]])
                else:
                    resp = fr24_call(
                        client.flight_summary.get_full,
                        callsigns=[item["callsign"]],
                        flight_datetime_from=day,
                        flight_datetime_to=day + timedelta(hours=23, minutes=59, seconds=59),
                    )
                matches = [m.model_dump() for m in (resp.data or [])]
                time.sleep(FR24_SLEEP_S)
                if not matches:
                    raise LookupError("no FR24 flight found")
                if len(matches) > 1:
                    log.warning("%s: %d FR24 candidates %s - picking best; pin one with an fr24_id column",
                                key, len(matches), [m["fr24_id"] for m in matches])
                summary = pick_best(matches, item)
                if not summary.get("hex"):
                    raise LookupError("FR24 summary has no hex/icao24")

                for label, got, want in (("reg", summary.get("reg"), item["registration"]),
                                         ("departure", summary.get("orig_icao"), item["departure"]),
                                         ("destination", summary.get("dest_icao"), item["destination"])):
                    if got and want and got.upper() != want:
                        log.warning("%s: sheet %s=%s differs from FR24 %s", key, label, want, got)

                tracks_resp = fr24_call(client.flight_tracks.get, summary["fr24_id"])
                time.sleep(FR24_SLEEP_S)
                tracks = []
                for ft in (tracks_resp.data or []):
                    tracks.extend(p.model_dump() for p in (ft.tracks or []))
                if len(tracks) < 2:
                    raise LookupError(f"FR24 track has {len(tracks)} point(s); need at least 2")

                out.append({"input": item, "key": key, "summary": summary, "tracks": tracks,
                            "fetched_at": datetime.now(timezone.utc).isoformat()})
                log.info("%s: fr24_id=%s hex=%s %s->%s, %d track points", key, summary["fr24_id"],
                         summary["hex"], summary.get("orig_icao"), summary.get("dest_icao"), len(tracks))
            except Exception as exc:
                failures += 1
                log.error("%s: FETCH FAILED - %s", key, exc)
                out.append({"input": item, "key": key, "error": str(exc)})
    save_json(args.state_dir / "fr24_flights.json", out)
    log.info("fetch done: %d ok, %d failed", len(out) - failures, failures)
    return 1 if failures == len(out) else 0


# --------------------------------------------------------------------------
# 2. adsb
# --------------------------------------------------------------------------

def find_existing(conn, icao24, fr24_id):
    """flight_sector (and its track) already created for this FR24 flight, if any."""
    return conn.execute(
        """select fs.flight_sector_id, ts.track_id
           from flight_tracking.flight_sector fs
           join flight_tracking.track_segment ts on ts.segment_id = fs.flight_sector_id
           where fs.icao24 = %s and fs.fr24_response_data->>'fr24_id' = %s
           order by fs.flight_sector_id desc limit 1""",
        (icao24, fr24_id)).fetchone()


PRESERVED_COLS = ("has_synched", "modification_lock_status", "mtow", "mtow_unit", "public_link",
                  "pod_link_sent", "pod_updated", "pod_response_data", "pod_updated_at", "pod_retry_count")


STANDALONE_SECTOR_SQL = """
insert into flight_tracking.flight_sector (
    flight_sector_id, icao24, last_segment_id, flight_start_e, flight_end_e, mvt_type, callsign, mode_a,
    mode_a_category, origin, destination, operator, aircraft_reg, aircraft_type, route_source,
    initial_airport_icao, initial_bay_nb, out_of_blocks_e, takeoff_airport, takeoff_time_e, takeoff_source,
    landing_airport, landing_time_e, landing_source, final_airport_icao, final_bay_nb, in_blocks_e, flight_state)
select ts.segment_id, ts.icao24, ts.segment_id, ts.segment_start_e, ts.segment_end_e,
    case when coalesce(air.n, 0) = 0 then 'SFC' else 'AIR' end,
    ts.callsign, ts.mode_a,
    case when ts.mode_a >= '4100' and ts.mode_a <= '4177' then 'DOMESTIC'::flight_tracking.FT_modea_category
         when ts.mode_a >= '4700' and ts.mode_a <= '4777' then 'DEPARTURE'::flight_tracking.FT_modea_category
         when ts.mode_a = '2000' then 'CONSPICUITY'::flight_tracking.FT_modea_category
         else 'INTERNATIONAL'::flight_tracking.FT_modea_category end,
    rt.origin, rt.destination, ac.operator, ac.aircraft_reg, ac.aircraft_type, rt.route_source,
    ts.initial_airport_icao, ts.initial_bay_nb, ts.out_of_blocks_e,
    tk.takeoff_airport, tk.takeoff_time_e, tk.takeoff_source,
    ld.landing_airport, ld.landing_time_e, ld.landing_source,
    ts.final_airport_icao, ts.final_bay_nb, ts.in_blocks_e, ts.segment_state
from flight_tracking.track_segment ts
left join lateral (select count(*) n from flight_tracking.track_event e
    where e.icao24 = ts.icao24 and e.event_type = 'AIR'
      and e.event_begin >= ts.segment_start_e and e.event_begin < ts.segment_end_e) air on true
left join lateral (select origin, destination, 'SENSOR'::flight_tracking.FT_event_source as route_source
    from flight_tracking.svl_route_raw rr
    where rr.callsign = ts.callsign and earliest_utc <= to_timestamp(ts.segment_end_e)
      and latest_utc >= to_timestamp(ts.segment_start_e)
    order by latest_utc desc limit 1) rt on true
left join lateral (select operator, aircraft_registration as aircraft_reg, aircraft_type
    from flight_tracking.svl_aircraft_raw ar
    where ar.hex_address = ts.icao24 and ar.effective_date >= ts.track_start and ar.effective_date <= ts.track_end
    order by ar.effective_date desc limit 1) ac on true
left join lateral (select airport_icao as takeoff_airport, event_begin as takeoff_time_e,
        'SENSOR'::flight_tracking.FT_event_source as takeoff_source
    from flight_tracking.track_event te
    where te.icao24 = ts.icao24 and te.event_type = 'TAKEOFF'
      and te.event_begin >= ts.segment_start_e and te.event_begin <= ts.segment_end_e
    order by event_begin limit 1) tk on true
left join lateral (select airport_icao as landing_airport, event_begin as landing_time_e,
        'SENSOR'::flight_tracking.FT_event_source as landing_source
    from flight_tracking.track_event le
    where le.icao24 = ts.icao24 and le.event_type = 'LANDING'
      and le.event_begin >= ts.segment_start_e and le.event_begin <= ts.segment_end_e
    order by event_begin desc limit 1) ld on true
where ts.segment_id = any(%s::bigint[])
on conflict (icao24, flight_start_e) do update set
    flight_end_e = excluded.flight_end_e, mvt_type = excluded.mvt_type, callsign = excluded.callsign,
    mode_a = excluded.mode_a, mode_a_category = excluded.mode_a_category, operator = excluded.operator,
    aircraft_reg = excluded.aircraft_reg, aircraft_type = excluded.aircraft_type,
    initial_airport_icao = excluded.initial_airport_icao, initial_bay_nb = excluded.initial_bay_nb,
    out_of_blocks_e = excluded.out_of_blocks_e, takeoff_airport = excluded.takeoff_airport,
    takeoff_time_e = excluded.takeoff_time_e, takeoff_source = excluded.takeoff_source,
    landing_airport = excluded.landing_airport, landing_time_e = excluded.landing_time_e,
    landing_source = excluded.landing_source, final_airport_icao = excluded.final_airport_icao,
    final_bay_nb = excluded.final_bay_nb, in_blocks_e = excluded.in_blocks_e, flight_state = excluded.flight_state
returning flight_sector_id
"""


def create_standalone_sectors(cur, segments):
    """Derive one flight_sector per segment, treating each backfilled segment as its own flight.

    ADSB's FT_flight_sector_init walks prior_segment_id links to find a flight's start/end, which
    yields nothing when the aircraft already has *later* segments (their prior_segment_id still
    points past this one). Backfilled flights don't need the link, so it is cleared and the same
    sector columns are derived from the segment alone; the FIR entry/exit and takeoff/landing
    estimation procedures then run on the new sector as usual. Existing segments are not touched.
    """
    cur.execute("update flight_tracking.track_segment set prior_segment_id = null, link_category = 'NEW' "
                "where segment_id = any(%s::bigint[])", (segments,))
    cur.execute(STANDALONE_SECTOR_SQL, (segments,))
    sectors = [r[0] for r in cur.fetchall()]
    if sectors:
        cur.execute("call flight_tracking.FT_flight_sector_update_fir_events(%s::bigint[])", (sectors,))
        cur.execute("call flight_tracking.FT_flight_sector_tko_est_update(%s::bigint[])", (sectors,))
    return sectors


def insert_one(conn, rec, existing=None):
    """Insert (or, when `existing=(sector_id, track_id)`, re-path) one FR24 flight in ADSB.

    An existing flight keeps its flight_sector_id (Billing's link): the track path is replaced
    and the sector is re-derived from the new path. Runs inside the caller's transaction."""
    s = rec["summary"]
    inp = rec["input"]
    icao24 = s["hex"].lower()
    pts = build_path_points(rec["tracks"])
    if len(pts) < 2:
        raise ValueError(f"only {len(pts)} of {len(rec['tracks'])} FR24 track points are inside the South Sudan FIR")
    track_start = datetime.fromtimestamp(pts[0][3], timezone.utc)
    track_end = datetime.fromtimestamp(pts[-1][3], timezone.utc)
    if (track_end - track_start).total_seconds() <= 30:
        raise ValueError("track is <= 30s long; ADSB post-processing ignores such tracks")

    cur = conn.cursor()
    saved, old_n = None, None
    if existing:
        sector_id0, track_id = existing
        cur.execute("select track_start, track_end, st_npoints(path) from flight_tracking.svl_track_raw where id = %s",
                    (track_id,))
        old_start, old_end, old_n = cur.fetchone()
        cur.execute(
            """update flight_tracking.svl_track_raw
               set path = %s::geometry, track_start = %s, track_end = %s where id = %s""",
            (path_ewkt(pts), track_start, track_end, track_id))
        # FT_track_segment_init only refreshes a segment's *end* fields (live tracks never move
        # their start), so bring all of the segment's timing fields in line with the new path.
        cur.execute(
            """update flight_tracking.track_segment es
               set track_start = su.track_start, track_end = su.track_end,
                   segment_start_e = su.segment_start_e, segment_end_e = su.segment_end_e,
                   trajectory_start_e = su.trajectory_start_e, trajectory_end_e = su.trajectory_end_e,
                   segment_state = su.segment_state
               from flight_tracking.FT_track_split(%s::bigint[]) su
               where es.track_id = su.track_id and es.split_id = su.split_id""", ([track_id],))
        if old_end != track_end:  # flight_raw marker sits at track_end
            cur.execute("delete from flight_tracking.svl_flight_raw where icao24 = %s and rx_utc = %s and sensor = 'FR24'",
                        (icao24, old_end))
        # The sector upsert is keyed on (icao24, flight_start_e); if the start moved, drop only the
        # sector row (the segment, hence the id, is kept) and restore Billing-related columns after.
        cur.execute(f"select {', '.join(PRESERVED_COLS)} from flight_tracking.flight_sector where flight_sector_id = %s",
                    (sector_id0,))
        saved = cur.fetchone()
        cur.execute("delete from flight_tracking.flight_sector where flight_sector_id = %s", (sector_id0,))
    else:
        cur.execute(
            """insert into flight_tracking.svl_track_raw(icao24, track_start, track_end, trk_state, path)
               values (%s, %s, %s, 'END'::flight_tracking.ATM_segment_state, %s::geometry)
               on conflict (icao24, track_start) do update
                 set track_end = excluded.track_end, trk_state = excluded.trk_state, path = excluded.path
               returning id""",
            (icao24, track_start, track_end, path_ewkt(pts)),
        )
        track_id = cur.fetchone()[0]

    # FT_compile_flights only considers END tracks that have a flight_raw row at track_end.
    squawk = (rec["tracks"][-1].get("squawk") or "")[:4] or None
    cur.execute(
        """insert into flight_tracking.svl_flight_raw(icao24, rx_utc, sensor, trust_lvl, trk_state, callsign, mode_a)
           values (%s, %s, 'FR24', 1, 'END'::flight_tracking.flight_data_state, %s, %s)
           on conflict (icao24, rx_utc) do nothing""",
        (icao24, track_end, (s.get("callsign") or rec["input"]["callsign"])[:8], squawk),
    )

    # Same chain FT_compile_flights runs, scoped to just this track.
    cur.execute("call flight_tracking.FT_get_track_events(%s::bigint[])", ([track_id],))
    if existing:  # segment already refreshed above; run the rest of FT_track_segment_process on it
        segments = [existing[0]]
        for proc in ("gate_usage", "profile", "csn", "mode_a", "link"):
            cur.execute(f"call flight_tracking.FT_track_segment_update_{proc}(%s::bigint[])", (segments,))
    else:
        cur.execute("select flight_tracking.FT_track_segment_process(%s::bigint[])", ([track_id],))
        segments = cur.fetchone()[0]
    if not segments:
        raise RuntimeError("ADSB produced no track_segment for this track")
    sectors = create_standalone_sectors(cur, segments)
    if not sectors:
        raise RuntimeError("no flight_sector derived for this track")

    # Choose the sector for this aircraft that covers the track.
    cur.execute(
        """select flight_sector_id from flight_tracking.flight_sector
           where flight_sector_id = any(%s::bigint[]) and icao24 = %s
             and flight_start_e <= %s and flight_end_e >= %s
           order by flight_start_e desc limit 1""",
        (sectors, icao24, pts[0][3] + 1, pts[-1][3] - 1),
    )
    row = cur.fetchone()
    if not row:
        raise RuntimeError(f"no derived flight_sector covers the track (sectors={sectors})")
    sector_id = row[0]
    if existing and sector_id != existing[0]:
        raise RuntimeError(f"flight_sector_id changed {existing[0]} -> {sector_id}; refusing (would break the Billing link)")
    if saved:
        cur.execute(f"update flight_tracking.flight_sector set {', '.join(c + ' = %s' for c in PRESERVED_COLS)} "
                    "where flight_sector_id = %s", (*saved, sector_id))

    # Overlay FR24 data. FR24 is the authority for a backfill, so non-null FR24 values win.
    from psycopg.types.json import Json
    cur.execute(
        """update flight_tracking.flight_sector set
               origin = coalesce(%(orig)s, origin)
              ,destination = coalesce(%(dest)s, destination)
              ,route_source = case when %(orig)s is not null or %(dest)s is not null
                                   then 'SENSOR'::flight_tracking.FT_event_source else route_source end
              ,aircraft_reg = coalesce(%(reg)s, aircraft_reg)
              ,aircraft_type = coalesce(%(type)s, aircraft_type)
              ,operator = coalesce(%(operator)s, operator)
              ,fr24_call = true
              ,fr24_response_data = %(resp)s
              ,fr24_enriched_at = now()
           where flight_sector_id = %(id)s""",
        {"orig": s.get("orig_icao") or inp.get("departure"), "dest": s.get("dest_icao") or inp.get("destination"),
         "reg": s.get("reg") or inp.get("registration"),
         "type": s.get("type") or inp.get("type"), "operator": s.get("operating_as") or inp.get("operator"),
         "resp": Json(s), "id": sector_id},
    )

    # Mark the track processed so the live pipeline does not pick it up again.
    cur.execute(
        """insert into flight_tracking.process_control(track_id, last_processed_time, complete)
           values (%s, %s, true)
           on conflict (track_id) do update set last_processed_time = excluded.last_processed_time, complete = true""",
        (track_id, track_end),
    )
    return {"track_id": track_id, "segment_ids": segments, "flight_sector_id": sector_id,
            "icao24": icao24, "path_points": len(pts),
            "action": f"updated path ({old_n} -> {len(pts)} points)" if existing else "created"}


def cmd_adsb(args):
    if not getattr(args, "skip_checks", False) and not run_checks("adsb"):
        return 1
    fetched = load_json(args.state_dir / "fr24_flights.json")
    results, failures = [], 0
    with pg_connect() as conn:
        for rec in fetched:
            if "error" in rec:
                log.warning("%s: skipped (fetch failed: %s)", rec["key"], rec["error"])
                continue
            try:
                with conn.transaction():
                    existing = find_existing(conn, rec["summary"]["hex"].lower(), rec["summary"]["fr24_id"])
                    info = insert_one(conn, rec, existing)
                    if args.dry_run:
                        raise _DryRunRollback(info)
                results.append({"key": rec["key"], "fr24_id": rec["summary"]["fr24_id"], **info})
                log.info("%s: %s flight_sector_id=%s track_id=%s (%d points)", rec["key"], info["action"],
                         info["flight_sector_id"], info["track_id"], info["path_points"])
            except _DryRunRollback as dr:
                log.info("%s: DRY RUN ok, would have %s flight_sector_id=%s (rolled back)",
                         rec["key"], dr.info["action"], dr.info["flight_sector_id"])
            except Exception as exc:
                failures += 1
                log.error("%s: ADSB INSERT FAILED (rolled back) - %s", rec["key"], exc)
                results.append({"key": rec["key"], "error": str(exc)})
    if not args.dry_run:
        save_json(args.state_dir / "adsb_flights.json", results)
    log.info("adsb done: %d ok, %d failed", len([r for r in results if "error" not in r]), failures)
    return 1 if failures else 0


class _DryRunRollback(Exception):
    def __init__(self, info):
        super().__init__("dry run")
        self.info = info


# --------------------------------------------------------------------------
# 3. billing
# --------------------------------------------------------------------------

def billing_token(root):
    r = requests.post(f"{root}/api/third-party/token/",
                      json={"token": env("BILLING_API_STATIC_TOKEN", required=True)}, timeout=15)
    r.raise_for_status()
    return r.json()["data"]["result"]["access"]


def load_payloads(state_dir):
    adsb = load_json(state_dir / "adsb_flights.json")
    payloads = []
    with pg_connect() as conn:
        for rec in adsb:
            if "error" in rec:
                continue
            row = conn.execute(FLIGHT_PAYLOAD_SQL, (rec["flight_sector_id"],)).fetchone()
            if not row:
                log.error("%s: flight_sector_id=%s vanished from ADSB", rec["key"], rec["flight_sector_id"])
                continue
            payloads.append((rec["key"], row[0]))
    return payloads


def _billing_counts(resp):
    """Counts live under data.result in the live API; tolerate data directly too."""
    d = (resp or {}).get("data") or {}
    return d.get("result") or d


def cmd_billing(args):
    # billing reads the flight_sector rows from ADSB, then posts them to Billing
    if not getattr(args, "skip_checks", False) and not run_checks("adsb", *([] if args.dry_run else ["billing"])):
        return 1
    payloads = load_payloads(args.state_dir)
    if not payloads:
        sys.exit("nothing to send")
    body = [json.loads(json.dumps(p, default=str)) for _, p in payloads]
    if args.dry_run:
        print(json.dumps(body, indent=2, default=str))
        log.info("dry run: %d payload(s) printed, nothing sent", len(body))
        return 0

    root = env("BILLING_API_ROOT", required=True).rstrip("/")
    headers = {"Authorization": f"Bearer {billing_token(root)}", "Content-Type": "application/json"}
    resp = requests.post(f"{root}/api/third-party-flight-dump/", headers=headers, json=body, timeout=60)
    try:
        data = resp.json()
    except ValueError:
        data = {"raw": resp.text[:2000]}
    save_json(args.state_dir / "billing_result.json",
              {"http_status": resp.status_code, "keys": [k for k, _ in payloads], "response": data})
    if resp.status_code != 200:
        log.error("Billing returned HTTP %s: %s", resp.status_code, json.dumps(data)[:1000])
        return 1
    counts = _billing_counts(data)
    log.info("billing: created=%s updated=%s failed=%s", counts.get("created_count"),
             counts.get("updated_count"), counts.get("failed_count"))
    for f in counts.get("failed_data") or []:
        log.error("billing rejected %s: %s", f.get("flight_identifier"), f.get("errors") or f.get("error"))
    return 1 if counts.get("failed_count") else 0


# --------------------------------------------------------------------------
# 4. verify (read-only)
# --------------------------------------------------------------------------

def cmd_verify(args):
    fetched = load_json(args.state_dir / "fr24_flights.json")
    adsb_path = args.state_dir / "adsb_flights.json"
    adsb = {r["key"]: r for r in load_json(adsb_path)} if adsb_path.exists() else {}
    bill_path = args.state_dir / "billing_result.json"
    billing = json.loads(bill_path.read_text()) if bill_path.exists() else None

    problems = 0
    conn = pg_connect() if adsb else None
    print(f"{'flight':<22}{'FR24':<8}{'ADSB':<8}{'BILLING':<9}notes")
    for rec in fetched:
        key, notes = rec["key"], []
        if "error" in rec:
            print(f"{key:<22}{'FAIL':<8}{'-':<8}{'-':<9}{rec['error']}")
            problems += 1
            continue
        s = rec["summary"]
        fr24_state = "ok"
        adsb_state = "-"
        if conn:
            a = adsb.get(key)
            if not a or "error" in a:
                adsb_state = "MISSING"
                notes.append((a or {}).get("error", "not inserted"))
                problems += 1
            else:
                row = conn.execute(
                    """select fs.icao24, fs.callsign, fs.origin, fs.destination, fs.aircraft_reg, fs.aircraft_type,
                              fs.fr24_call, fs.flight_state::text, st_npoints(tr.path), fs.fir_entry_time_e is not null
                       from flight_tracking.flight_sector fs
                       left join flight_tracking.svl_track_raw tr on tr.id = %s
                       where fs.flight_sector_id = %s""",
                    (a["track_id"], a["flight_sector_id"])).fetchone()
                if not row:
                    adsb_state = "MISSING"
                    notes.append(f"flight_sector {a['flight_sector_id']} not found")
                    problems += 1
                else:
                    icao24, cs, org, dst, reg, typ, fr24_call, fstate, npts, has_fir = row
                    adsb_state = "ok"
                    checks = [("icao24", icao24, (s.get("hex") or "").lower()),
                              ("origin", org, s.get("orig_icao")), ("dest", dst, s.get("dest_icao")),
                              ("reg", reg, s.get("reg")), ("type", typ, s.get("type"))]
                    for name, got, want in checks:
                        if want and (got or "").strip().lower() != want.strip().lower():
                            notes.append(f"{name}: adsb={got!r} fr24={want!r}")
                    if npts != len(build_path_points(rec["tracks"])):
                        notes.append(f"path points adsb={npts} fr24={len(build_path_points(rec['tracks']))}")
                    if not fr24_call:
                        notes.append("fr24_call=false")
                    if not has_fir:
                        notes.append("no FIR entry derived")
                    if notes:
                        adsb_state = "DIFF"
                        problems += 1
                    notes.append(f"sector={a['flight_sector_id']} state={fstate}")
        bill_state = "-"
        if billing:
            bill_state = "sent" if key in billing.get("keys", []) else "NOT SENT"
            if billing["http_status"] != 200:
                bill_state = f"HTTP {billing['http_status']}"
            for f in (_billing_counts(billing["response"]).get("failed_data") or []):
                sid = adsb.get(key, {}).get("flight_sector_id")
                if sid and str(sid) in str(f.get("flight_identifier")):
                    bill_state = "REJECTED"
                    notes.append(str(f.get("errors") or f.get("error")))
                    problems += 1
        print(f"{key:<22}{fr24_state:<8}{adsb_state:<8}{bill_state:<9}{'; '.join(notes)}")
    if conn:
        conn.close()
    print(f"\n{problems} problem(s).")
    return 1 if problems else 0


def cmd_all(args):
    """check -> fetch -> adsb -> billing -> verify, stopping if a stage cannot continue."""
    if not run_checks("adsb", "billing"):  # before any FR24 credits are spent
        return 1
    args.skip_checks = True
    log.info("=== [1/4] fetch")
    if cmd_fetch(args) != 0:
        log.error("cancelled: fetch produced no usable flights")
        return 1
    log.info("=== [2/4] adsb%s", " (dry run)" if args.dry_run else "")
    rc = cmd_adsb(args)
    if args.dry_run:
        log.info("dry run finished: billing and verify skipped (nothing was written)")
        return rc
    done = [r for r in load_json(args.state_dir / "adsb_flights.json") if "error" not in r]
    if not done:
        log.error("cancelled: no flight made it into ADSB, nothing to send to Billing")
        return 1
    if rc:
        log.warning("some flights failed in ADSB; continuing with the %d that succeeded", len(done))
    log.info("=== [3/4] billing")
    if cmd_billing(args) != 0:
        log.error("billing step reported failures; running verify anyway")
        cmd_verify(args)
        return 1
    log.info("=== [4/4] verify")
    return cmd_verify(args)


# --------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--state-dir", type=Path, default=Path("state"), help="where stage output files live")
    ap.add_argument("--env-file", default=".env")
    sub = ap.add_subparsers(dest="command", required=True)

    p = sub.add_parser("fetch", help="fetch flight details + track from FR24")
    p.add_argument("input", type=Path, help=".xlsx/.csv/.tsv flight list (CALL SIGN column required)")
    p.add_argument("--date", metavar="YYYY-MM-DD", help="operation date for rows without a date column (required unless the sheet has one)")
    p.add_argument("--sheet", help="workbook tab to read (default: active tab)")
    p.set_defaults(fn=cmd_fetch)

    p = sub.add_parser("adsb", help="insert fetched flights into the ADSB database")
    p.add_argument("--dry-run", action="store_true", help="run everything, then roll back")
    p.set_defaults(fn=cmd_adsb)

    p = sub.add_parser("billing", help="push the ADSB flight_sector rows to Billing")
    p.add_argument("--dry-run", action="store_true", help="print the payload, send nothing")
    p.set_defaults(fn=cmd_billing)

    p = sub.add_parser("check", help="test the ADSB database and Billing endpoint connections")
    p.set_defaults(fn=cmd_check)

    p = sub.add_parser("all", help="check, fetch, adsb, billing and verify in sequence")
    p.add_argument("input", type=Path, help=".xlsx/.csv/.tsv flight list (CALL SIGN column required)")
    p.add_argument("--sheet", help="workbook tab to read (default: active tab)")
    p.add_argument("--date", metavar="YYYY-MM-DD", help="operation date for rows without a date column (required unless the sheet has one)")
    p.add_argument("--dry-run", action="store_true", help="checks + fetch + ADSB dry run only; nothing is written or sent")
    p.set_defaults(fn=cmd_all)

    p = sub.add_parser("verify", help="read-only check of all stages")
    p.set_defaults(fn=cmd_verify)

    args = ap.parse_args()
    load_dotenv(args.env_file)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    sys.exit(args.fn(args))


if __name__ == "__main__":
    main()
